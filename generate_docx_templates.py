import os
import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

os.makedirs("templates", exist_ok=True)

def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    tcPr.append(parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>'))

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def add_header_bps(doc, title_sub=""):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r1 = p.add_run("BADAN PUSAT STATISTIK KABUPATEN KEPULAUAN SELAYAR\n")
    r1.bold = True
    r1.font.size = Pt(13)
    r1.font.name = "Arial"
    r1.font.color.rgb = RGBColor(26, 54, 93)

    r2 = p.add_run("Jalan Ahmad Yani No. 12, Benteng, Kabupaten Kepulauan Selayar, Sulawesi Selatan 92812\nTelepon: (0414) 21102 | Email: bps7301@bps.go.id | Website: selayarkab.bps.go.id\n")
    r2.font.size = Pt(8.5)
    r2.font.italic = True
    r2.font.color.rgb = RGBColor(100, 116, 139)

    # Line divider
    p_line = doc.add_paragraph()
    p_line.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_line = p_line.add_run("―" * 55)
    r_line.font.bold = True
    r_line.font.color.rgb = RGBColor(198, 149, 59)
    p_line.paragraph_format.space_after = Pt(12)

# 1. SK Tim Pembangunan ZI
def make_sk_tim_zi():
    doc = Document()
    add_header_bps(doc)

    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_t1 = p_title.add_run("KEPUTUSAN KEPALA BADAN PUSAT STATISTIK\nKABUPATEN KEPULAUAN SELAYAR\n")
    r_t1.bold = True
    r_t1.font.size = Pt(12)
    r_t2 = p_title.add_run("NOMOR: B-012/73010/OT.000/01/2025\n\nTENTANG\n")
    r_t2.bold = True
    r_t2.font.size = Pt(11)
    r_t3 = p_title.add_run("PEMBENTUKAN TIM KERJA PEMBANGUNAN ZONA INTEGRITAS (ZI)\nMENUJU WILAYAH BEBAS DARI KORUPSI (WBK) DAN WILAYAH BIROKRASI BERSIH DAN MELAYANI (WBBM)\nBPS KABUPATEN KEPULAUAN SELAYAR TAHUN 2025")
    r_t3.bold = True
    r_t3.font.size = Pt(11)
    p_title.paragraph_format.space_after = Pt(14)

    p_menimbang = doc.add_paragraph()
    r = p_menimbang.add_run("Menimbang:\n")
    r.bold = True
    p_menimbang.add_run("a. bahwa dalam rangka mengakselerasi pencapaian tujuan Reformasi Birokrasi di lingkungan BPS Kabupaten Kepulauan Selayar, perlu dilakukan pembangunan Zona Integritas menuju WBK/WBBM;\n"
                        "b. bahwa untuk menjamin kelancaran, keterpaduan, dan akuntabilitas pelaksanaan pembangunan Zona Integritas, perlu dibentuk Tim Kerja Pembangunan Zona Integritas;\n"
                        "c. bahwa berdasarkan pertimbangan sebagaimana dimaksud pada huruf a dan huruf b, perlu menetapkan Keputusan Kepala Badan Pusat Statistik Kabupaten Kepulauan Selayar.")

    p_mengingat = doc.add_paragraph()
    r = p_mengingat.add_run("Mengingat:\n")
    r.bold = True
    p_mengingat.add_run("1. Undang-Undang Nomor 16 Tahun 1997 tentang Statistik;\n"
                        "2. Peraturan Presiden Nomor 81 Tahun 2010 tentang Grand Design Reformasi Birokrasi 2010-2025;\n"
                        "3. Peraturan Menteri PANRB Nomor 90 Tahun 2021 tentang Pembangunan dan Evaluasi Zona Integritas Menuju WBK/WBBM;\n"
                        "4. Peraturan BPS Nomor 5 Tahun 2023 tentang Organisasi dan Tata Kerja BPS Provinsi dan BPS Kabupaten/Kota.")

    p_putus = doc.add_paragraph()
    p_putus.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_putus = p_putus.add_run("MEMUTUSKAN:")
    r_putus.bold = True
    p_putus.paragraph_format.space_before = Pt(8)
    p_putus.paragraph_format.space_after = Pt(8)

    doc.add_paragraph("Menetapkan: KEPUTUSAN KEPALA BADAN PUSAT STATISTIK KABUPATEN KEPULAUAN SELAYAR TENTANG PEMBENTUKAN TIM KERJA PEMBANGUNAN ZONA INTEGRITAS MENUJU WBK/WBBM TAHUN 2025.")
    doc.add_paragraph("KESATU: Membentuk Tim Kerja Pembangunan Zona Integritas Menuju WBK/WBBM BPS Kabupaten Kepulauan Selayar Tahun 2025 dengan susunan keanggotaan sebagaimana tercantum dalam Lampiran yang merupakan bagian tidak terpisahkan dari Keputusan ini.")
    doc.add_paragraph("KEDUA: Tim Kerja sebagaimana dimaksud Diktum KESATU bertugas menyusun rencana kerja, mengoordinasikan pengumpulan bukti dukung LKE pada 6 (enam) Area Perubahan, memonitor implementasi, dan menyusun laporan tindak lanjut.")
    doc.add_paragraph("KETIGA: Segala biaya yang timbul sebagai akibat ditetapkannya Keputusan ini dibebankan pada DIPA BPS Kabupaten Kepulauan Selayar Tahun Anggaran 2025.")
    doc.add_paragraph("KEEMPAT: Keputusan ini mulai berlaku pada tanggal ditetapkan dengan ketentuan apabila di kemudian hari terdapat kekeliruan akan dilakukan perbaikan sebagaimana mestinya.")

    # Tanda tangan
    p_ttd = doc.add_paragraph()
    p_ttd.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p_ttd.paragraph_format.space_before = Pt(20)
    p_ttd.add_run("Ditetapkan di Benteng\npada tanggal 02 Januari 2025\n\nKepala BPS Kabupaten Kepulauan Selayar,\n\n\n\n")
    r_nama = p_ttd.add_run("(Nama Kepala BPS Kab. Kepulauan Selayar)\nNIP. 19xxxxxxxxxxxxxx")
    r_nama.bold = True

    # Page break for attachment
    doc.add_page_break()
    p_lamp = doc.add_paragraph()
    p_lamp.add_run("LAMPIRAN\nKEPUTUSAN KEPALA BPS KABUPATEN KEPULAUAN SELAYAR\nNOMOR: B-012/73010/OT.000/01/2025\nTENTANG SUSUNAN TIM KERJA PEMBANGUNAN ZONA INTEGRITAS TAHUN 2025\n").bold = True

    table = doc.add_table(rows=1, cols=4)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    hdr = table.rows[0].cells
    hdr[0].text = "No"
    hdr[1].text = "Area Perubahan"
    hdr[2].text = "Kedudukan dalam Tim"
    hdr[3].text = "Nama / Jabatan Pegawai"
    for c in hdr:
        set_cell_background(c, "1A365D")
        for p in c.paragraphs:
            p.runs[0].font.bold = True
            p.runs[0].font.color.rgb = RGBColor(255, 255, 255)

    data_tim = [
        ("1", "Pengarah", "Pembina / Penanggung Jawab", "Kepala BPS Kab. Kepulauan Selayar"),
        ("2", "Ketua Pembangunan ZI", "Ketua Tim Kerja", "Kasubbag Umum"),
        ("3", "Sekretariat", "Koordinator Data & Administrasi", "Pranata Komputer Ahli Muda"),
        ("4", "Area 1: Manajemen Perubahan", "Koordinator & Anggota", "Statistisi Ahli Pertama / Agen Perubahan"),
        ("5", "Area 2: Penataan Tatalaksana", "Koordinator & Anggota", "Pranata Komputer / Pengelola SPBE"),
        ("6", "Area 3: Penataan Manajemen SDM", "Koordinator & Anggota", "Analis Kepegawaian / Pengelola SDM"),
        ("7", "Area 4: Penguatan Akuntabilitas", "Koordinator & Anggota", "Penyusun LAKIN / Pengelola SAKIP"),
        ("8", "Area 5: Penguatan Pengawasan", "Koordinator & Anggota", "Petugas UPG / WBS / SPIP"),
        ("9", "Area 6: Pelayanan Publik", "Koordinator & Anggota", "Koordinator Pelayanan Statistik Terpadu (PST)")
    ]

    for row_data in data_tim:
        row = table.add_row().cells
        for idx, text in enumerate(row_data):
            row[idx].text = text
            set_cell_margins(row[idx], 80, 80, 100, 100)

    doc.save("templates/Format_01_SK_Tim_Pembangunan_ZI_2025_BPS_Selayar.docx")

# 2. Notula Rapat Pembentukan Tim
def make_notula_pembentukan():
    doc = Document()
    add_header_bps(doc)

    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_t = p_title.add_run("NOTULA RAPAT PEMBENTUKAN TIM KERJA PEMBANGUNAN ZONA INTEGRITAS\nBPS KABUPATEN KEPULAUAN SELAYAR TAHUN 2025")
    r_t.bold = True
    r_t.font.size = Pt(12)

    doc.add_paragraph("Hari / Tanggal  : Jumat, 03 Januari 2025\nWaktu           : 09.00 - 11.30 WITA\nTempat          : Ruang Rapat BPS Kabupaten Kepulauan Selayar\nPemimpin Rapat  : Kepala BPS Kabupaten Kepulauan Selayar\nNotulis         : Sekretaris Tim ZI\nPeserta Rapat   : Seluruh Pegawai BPS Kabupaten Kepulauan Selayar (18 Pegawai)")

    doc.add_heading("I. Agenda Rapat", level=2)
    doc.add_paragraph("1. Arahan Kepala BPS Kab. Kepulauan Selayar terkait Komitmen Pembangunan Zona Integritas Menuju WBK Tahun 2025.\n"
                      "2. Pembahasan kriteria dan mekanisme pemilihan anggota Tim Kerja Pembangunan ZI.\n"
                      "3. Penentuan Koordinator dan Anggota pada 6 (enam) Area Perubahan.\n"
                      "4. Penyusunan timeline pengumpulan dan digitalisasi bukti dukung pada Google Drive.")

    doc.add_heading("II. Mekanisme Pemilihan Anggota Tim Kerja", level=2)
    doc.add_paragraph("Pemilihan anggota Tim Kerja dilakukan secara terbuka dan objektif dengan memperhatikan:\n"
                      "a. Kompetensi teknis sesuai bidang area pengungkit;\n"
                      "b. Komitmen integritas tinggi (tidak pernah dijatuhi hukuman disiplin);\n"
                      "c. Keterwakilan dari seluruh fungsi/unsur pada BPS Kabupaten Kepulauan Selayar;\n"
                      "d. Rekomendasi hasil musyawarah mufakat seluruh pegawai.")

    doc.add_heading("III. Hasil Kesepakatan Rapat", level=2)
    doc.add_paragraph("1. Disepakati struktur Tim Kerja ZI 2025 yang dipimpin oleh Kasubbag Umum.\n"
                      "2. Telah ditetapkan penanggung jawab (PIC) untuk masing-masing dari 6 Area Perubahan.\n"
                      "3. Semua bukti dukung wajib diarsipkan dalam Google Drive resmi dengan penamaan standar.\n"
                      "4. Rapat evaluasi berkala akan dilaksanakan setiap bulan (minggu ke-4).")

    p_ttd = doc.add_paragraph()
    p_ttd.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p_ttd.paragraph_format.space_before = Pt(24)
    p_ttd.add_run("Benteng, 03 Januari 2025\n\nMengetahui,\nKepala BPS Kabupaten Kepulauan Selayar\n\n\n\n(Nama Kepala BPS)\nNIP. 19xxxxxxxxxxxxxx")

    doc.save("templates/Format_02_Notula_Rapat_Pembentukan_Tim_ZI_BPS_Selayar.docx")

# 3. Surat Pernyataan Bebas Benturan Kepentingan
def make_surat_benturan_kepentingan():
    doc = Document()
    add_header_bps(doc)

    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_t = p_title.add_run("SURAT PERNYATAAN BEBAS BENTURAN KEPENTINGAN\nBPS KABUPATEN KEPULAUAN SELAYAR")
    r_t.bold = True
    r_t.font.size = Pt(12)
    p_title.paragraph_format.space_after = Pt(14)

    doc.add_paragraph("Yang bertanda tangan di bawah ini:")
    doc.add_paragraph("Nama               : .......................................................................\n"
                      "NIP                : .......................................................................\n"
                      "Pangkat / Golongan : .......................................................................\n"
                      "Jabatan            : .......................................................................\n"
                      "Unit Kerja         : BPS Kabupaten Kepulauan Selayar")

    doc.add_paragraph("Dengan ini menyatakan dengan sesungguhnya bahwa saya:")
    doc.add_paragraph("1. Dalam menjalankan tugas kedinasan dan wewenang yang diamanahkan kepada saya, senantiasa mematuhi peraturan perundang-undangan, kode etik, dan prinsip-prinsip tata kelola pemerintahan yang baik (good governance);\n"
                      "2. Tidak memiliki hubungan afiliasi atau kepentingan pribadi/keluarga/golongan yang dapat mempengaruhi independensi, netralitas, dan objektivitas saya dalam pengambilan keputusan kedinasan;\n"
                      "3. Bersedia melaporkan segera secara tertulis kepada Pimpinan Satker atau Tim Penanganan Benturan Kepentingan apabila dalam pelaksanaan tugas terdapat potensi atau situasi benturan kepentingan;\n"
                      "4. Bersedia menerima sanksi sesuai ketentuan perundang-undangan yang berlaku apabila pernyataan yang saya sampaikan ini terbukti tidak benar.")

    doc.add_paragraph("Demikian surat pernyataan ini saya buat dengan sadar dan penuh rasa tanggung jawab untuk dipergunakan sebagaimana mestinya.")

    p_ttd = doc.add_paragraph()
    p_ttd.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p_ttd.paragraph_format.space_before = Pt(20)
    p_ttd.add_run("Benteng, .............................. 2025\nYang Membuat Pernyataan,\n\n\n\nMaterai Rp 10.000\n\n\n\n(......................................................)\nNIP. ...............................................")

    doc.save("templates/Format_03_Surat_Pernyataan_Bebas_Benturan_Kepentingan_BPS_Selayar.docx")

# 4. Laporan Monev ZI Triwulan
def make_laporan_monev():
    doc = Document()
    add_header_bps(doc)

    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_t = p_title.add_run("LAPORAN MONITORING DAN EVALUASI TRIWULANAN\nPEMBANGUNAN ZONA INTEGRITAS MENUJU WBK/WBBM\nBPS KABUPATEN KEPULAUAN SELAYAR\nTRIWULAN I TAHUN 2025")
    r_t.bold = True
    r_t.font.size = Pt(12)

    doc.add_heading("I. Pendahuluan", level=2)
    doc.add_paragraph("Laporan Monitoring dan Evaluasi (Monev) ini disusun sebagai bentuk akuntabilitas dan kontrol terhadap implementasi rencana kerja Pembangunan Zona Integritas pada BPS Kabupaten Kepulauan Selayar selama periode Triwulan I Tahun 2025. Monitoring ini dipimpin langsung oleh Kepala BPS Kab. Kepulauan Selayar bersama seluruh anggota Tim Kerja ZI.")

    doc.add_heading("II. Capaian Pengumpulan Bukti Dukung per Area", level=2)
    table = doc.add_table(rows=1, cols=5)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    hdr = table.rows[0].cells
    hdr[0].text = "Area"
    hdr[1].text = "Nama Area Perubahan"
    hdr[2].text = "Target TW I"
    hdr[3].text = "Realisasi"
    hdr[4].text = "Status / Keterangan"
    for c in hdr:
        set_cell_background(c, "1A365D")
        for p in c.paragraphs:
            p.runs[0].font.bold = True
            p.runs[0].font.color.rgb = RGBColor(255, 255, 255)

    areas = [
        ("Area 1", "Manajemen Perubahan", "100%", "75%", "Perlu penguatan peran pimpinan sebagai role model"),
        ("Area 2", "Penataan Tatalaksana", "100%", "96%", "Optimal, evaluasi SOP terjadwal"),
        ("Area 3", "Penataan Manajemen SDM", "100%", "80%", "TNA terlaksana, dokumen riwayat diklat diproses"),
        ("Area 4", "Penguatan Akuntabilitas", "100%", "100%", "LAKIN selesai dan diunggah tepat waktu"),
        ("Area 5", "Penguatan Pengawasan", "100%", "65%", "Identifikasi benturan kepentingan perlu dilengkapi"),
        ("Area 6", "Peningkatan Pelayanan Publik", "100%", "85%", "Survei SKD TW I berjalan, maklumat terpasang")
    ]
    for a in areas:
        row = table.add_row().cells
        for idx, val in enumerate(a):
            row[idx].text = val
            set_cell_margins(row[idx], 60, 60, 80, 80)

    doc.add_heading("III. Analisis Kendala dan Solusi Tindak Lanjut", level=2)
    doc.add_paragraph("1. Kendala: Beberapa dokumen bukti dukung belum terstruktur di Google Drive.\n"
                      "   Solusi: Tim Sekretariat ZI melakukan pendampingan pengunggahan dan standardisasi penamaan file per sub-indikator.\n"
                      "2. Kendala: Sosialisasi benturan kepentingan belum menyeluruh kepada seluruh ASN dan PPNPN.\n"
                      "   Solusi: Mengagendakan internalisasi khusus pada apel pagi dan knowledge sharing mingguan.")

    p_ttd = doc.add_paragraph()
    p_ttd.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p_ttd.paragraph_format.space_before = Pt(20)
    p_ttd.add_run("Benteng, 31 Maret 2025\nKetua Tim Pembangunan ZI,\n\n\n\n(Kasubbag Umum BPS Kab. Kepulauan Selayar)\nNIP. 19xxxxxxxxxxxxxx")

    doc.save("templates/Format_04_Laporan_Monev_ZI_Triwulan_BPS_Selayar.docx")

# 5. Pakta Integritas Pegawai
def make_pakta_integritas():
    doc = Document()
    add_header_bps(doc)

    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_t = p_title.add_run("PAKTA INTEGRITAS PEGAWAI\nDALAM RANGKA PEMBANGUNAN ZONA INTEGRITAS MENUJU WBK/WBBM\nBPS KABUPATEN KEPULAUAN SELAYAR TAHUN 2025")
    r_t.bold = True
    r_t.font.size = Pt(12)

    doc.add_paragraph("Saya yang bertanda tangan di bawah ini:")
    doc.add_paragraph("Nama        : .......................................................................\n"
                      "NIP         : .......................................................................\n"
                      "Jabatan     : .......................................................................\n"
                      "Unit Kerja  : BPS Kabupaten Kepulauan Selayar")

    doc.add_paragraph("Dengan sungguh-sungguh menyatakan janji dan komitmen bahwa saya:")
    doc.add_paragraph("1. Berperan secara proaktif dalam upaya pencegahan dan pemberantasan Korupsi, Kolusi, dan Nepotisme serta tidak melibatkan diri dalam perbuatan tercela;\n"
                      "2. Tidak meminta atau menerima pemberian secara langsung atau tidak langsung berupa suap, hadiah, bantuan, atau bentuk lainnya yang tidak sesuai dengan ketentuan yang berlaku (bebas gratifikasi);\n"
                      "3. Bersikap transparan, jujur, objektif, dan akuntabel dalam melaksanakan tugas pelayanan statistik kepada seluruh pemangku kepentingan dan masyarakat umum;\n"
                      "4. Menjunjung tinggi Core Values ASN BerAKHLAK (Berorientasi Pelayanan, Akuntabel, Kompeten, Harmonis, Loyal, Adaptif, Kolaboratif) dalam kehidupan kerja sehari-hari;\n"
                      "5. Apabila saya melanggar hal-hal yang telah saya nyatakan dalam Pakta Integritas ini, saya bersedia dikenakan sanksi moral, sanksi administratif, serta dituntut ganti rugi dan pidana sesuai ketentuan perundang-undangan yang berlaku.")

    p_ttd = doc.add_paragraph()
    p_ttd.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p_ttd.paragraph_format.space_before = Pt(20)
    p_ttd.add_run("Benteng, 06 Januari 2025\nPegawai yang Menyatakan,\n\n\n\nMaterai Rp 10.000\n\n\n\n(......................................................)\nNIP. ...............................................")

    doc.save("templates/Format_05_Pakta_Integritas_Pegawai_BPS_Selayar.docx")

# 6. Pedoman Reward and Punishment Pelayanan
def make_reward_punishment():
    doc = Document()
    add_header_bps(doc)

    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_t = p_title.add_run("STANDAR OPERASIONAL PROSEDUR / KEBIJAKAN\nPEMBERIAN PENGHARGAAN (REWARD) DAN SANKSI (PUNISHMENT)\nBAGI PETUGAS PEMBERI PELAYANAN STATISTIK TERPADU (PST)\nBPS KABUPATEN KEPULAUAN SELAYAR")
    r_t.bold = True
    r_t.font.size = Pt(12)

    doc.add_heading("I. Tujuan", level=2)
    doc.add_paragraph("Mendorong peningkatan profesionalisme, integritas, dan keramahan petugas pelayanan PST BPS Kabupaten Kepulauan Selayar serta memberikan standar perlakuan yang adil terhadap kinerja prima maupun pelanggaran standar pelayanan.")

    doc.add_heading("II. Unsur Penilaian Petugas Pelayanan", level=2)
    doc.add_paragraph("Penilaian petugas PST dilakukan secara berkala tiap semester dengan 3 (tiga) komponen utama:\n"
                      "1. Disiplin Kehadiran (bobot 30%): Ketepatan waktu membuka loket dan kepatuhan jam piket layanan.\n"
                      "2. Kualitas Kinerja Layanan (bobot 35%): Kecepatan respon permintaan data dan kepatuhan SOP PST.\n"
                      "3. Kepuasan Pengguna Layanan (bobot 35%): Hasil rating Three Color Survey / SKM pengunjung PST.")

    doc.add_heading("III. Bentuk Penghargaan (Reward)", level=2)
    doc.add_paragraph("Bagi petugas dengan nilai tertinggi dan predikat 'Sangat Memuaskan':\n"
                      "- Piagam Penghargaan 'Petugas Layanan Teladan' (Employee of the Quarter);\n"
                      "- Prioritas pengikutsertaan dalam pelatihan teknis / capacity building;\n"
                      "- Pemberian cinderamata apresiasi satker.")

    doc.add_heading("IV. Bentuk Sanksi (Punishment)", level=2)
    doc.add_paragraph("Bagi petugas yang melakukan pelanggaran standar pelayanan atau memperoleh aduan yang terbukti valid:\n"
                      "- Teguran lisan dan bimbingan konseling oleh Pembina PST;\n"
                      "- Surat Peringatan Tertulis apabila tidak ada perbaikan dalam kurun waktu 14 hari kerja;\n"
                      "- Penarikan dari penugasan loket pelayanan publik ke tugas administrasi internal.")

    p_ttd = doc.add_paragraph()
    p_ttd.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p_ttd.paragraph_format.space_before = Pt(20)
    p_ttd.add_run("Ditetapkan di Benteng\npada tanggal 02 Januari 2025\n\nKepala BPS Kabupaten Kepulauan Selayar,\n\n\n\n(Nama Kepala BPS)\nNIP. 19xxxxxxxxxxxxxx")

    doc.save("templates/Format_06_Pedoman_Reward_Punishment_Petugas_Layanan_BPS_Selayar.docx")

# Run all generator functions
make_sk_tim_zi()
make_notula_pembentukan()
make_surat_benturan_kepentingan()
make_laporan_monev()
make_pakta_integritas()
make_reward_punishment()

print("All docx templates created successfully.")
