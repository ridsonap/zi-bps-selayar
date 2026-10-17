import os
import json
import shutil
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

with open("data.json", "r", encoding="utf-8") as f:
    app_data = json.load(f)

# Define target root directories
ROOT_LOCAL = os.path.abspath("Bukti_Dukung_Selayar_2025")
GDRIVE_BASE = os.path.expanduser("~/Library/CloudStorage/GoogleDrive-ridsonalfarizal15@gmail.com/My Drive")
ROOT_GDRIVE_ZI = os.path.join(GDRIVE_BASE, "ZI", "Bukti_Dukung_Selayar_2025")
ROOT_GDRIVE_SELAYAR = os.path.join(GDRIVE_BASE, "Selayar_ZI", "Bukti_Dukung_Selayar_2025")

targets = [ROOT_LOCAL]
if os.path.exists(os.path.join(GDRIVE_BASE, "ZI")):
    targets.append(ROOT_GDRIVE_ZI)
if os.path.exists(os.path.join(GDRIVE_BASE, "Selayar_ZI")):
    targets.append(ROOT_GDRIVE_SELAYAR)
else:
    # Also create in Selayar_ZI if parent can be made
    try:
        os.makedirs(os.path.join(GDRIVE_BASE, "Selayar_ZI"), exist_ok=True)
        targets.append(ROOT_GDRIVE_SELAYAR)
    except Exception as e:
        print("Note on Selayar_ZI:", e)

print("Target roots for creation:", targets)

def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(f'''
        <w:tcMar {nsdecls("w")}>
            <w:top w:w="{top}" w:type="dxa"/>
            <w:bottom w:w="{bottom}" w:type="dxa"/>
            <w:left w:w="{left}" w:type="dxa"/>
            <w:right w:w="{right}" w:type="dxa"/>
        </w:tcMar>
    ''')
    tcPr.append(tcMar)

def sanitize_folder_name(name):
    # Keep it clean for macOS/Windows/Google Drive
    name = name.replace(" / ", "_").replace("/", "_").replace("\\", "_")
    name = name.replace(":", "_").replace("*", "_").replace("?", "_")
    name = name.replace("\"", "_").replace("<", "_").replace(">", "_").replace("|", "_")
    name = " ".join(name.split())
    return name.strip()

def create_indicator_docx(filepath, item, area_num, area_title, aspect, group_name):
    doc = docx.Document()
    
    # Page setup - A4, 1 inch margins
    sections = doc.sections
    for section in sections:
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1)
        section.right_margin = Inches(1)
        section.page_width = Inches(8.27)
        section.page_height = Inches(11.69)

    # 1. Official Kop
    p_kop = doc.add_paragraph()
    p_kop.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_kop.paragraph_format.space_before = Pt(0)
    p_kop.paragraph_format.space_after = Pt(2)
    
    r1 = p_kop.add_run("BADAN PUSAT STATISTIK KABUPATEN KEPULAUAN SELAYAR\n")
    r1.bold = True
    r1.font.size = Pt(13)
    r1.font.name = "Arial"
    r1.font.color.rgb = RGBColor(30, 58, 138)
    
    r2 = p_kop.add_run("Jl. Ahmad Yani No. 12, Benteng, Kepulauan Selayar, Sulawesi Selatan 92812\n")
    r2.font.size = Pt(8.5)
    r2.font.name = "Arial"
    r2.font.color.rgb = RGBColor(71, 85, 105)
    
    r3 = p_kop.add_run("Telepon: (0414) 21102 | Pos-el: bps7301@bps.go.id | Laman: selayarkab.bps.go.id")
    r3.font.size = Pt(8.5)
    r3.font.name = "Arial"
    r3.font.color.rgb = RGBColor(71, 85, 105)

    # Divider line
    p_line = doc.add_paragraph()
    p_line.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_line.paragraph_format.space_before = Pt(2)
    p_line.paragraph_format.space_after = Pt(14)
    r_line = p_line.add_run("━" * 55)
    r_line.font.size = Pt(10)
    r_line.font.color.rgb = RGBColor(30, 58, 138)

    # Title Box
    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_title.paragraph_format.space_before = Pt(0)
    p_title.paragraph_format.space_after = Pt(2)
    r_title = p_title.add_run("PANDUAN & FORMAT BUKTI DUKUNG LEMBAR KERJA EVALUASI (LKE)")
    r_title.bold = True
    r_title.font.size = Pt(11.5)
    r_title.font.name = "Arial"

    p_sub = doc.add_paragraph()
    p_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_sub.paragraph_format.space_before = Pt(0)
    p_sub.paragraph_format.space_after = Pt(14)
    r_sub = p_sub.add_run(f"PEMBANGUNAN ZONA INTEGRITAS MENUJU WBK/WBBM TAHUN 2025")
    r_sub.font.size = Pt(10)
    r_sub.font.name = "Arial"
    r_sub.font.color.rgb = RGBColor(100, 116, 139)

    # Meta Table
    tbl_meta = doc.add_table(rows=6, cols=2)
    tbl_meta.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl_meta.autofit = False
    
    meta_rows = [
        ("Area Perubahan", f"Area {area_num}: {area_title}"),
        ("Komponen & Sub", f"{aspect} — {group_name}"),
        ("Kode Indikator", item["code"]),
        ("Indikator LKE", item["question"]),
        ("Target / Kriteria BPS", item.get("kriteria", "-")),
        ("Status Evaluasi Satker", f"Jawaban: {item.get('jawaban', '-')} | Nilai: {item.get('nilai', '0.00')} ({item.get('badge', '-')})")
    ]
    
    for idx, (label, val) in enumerate(meta_rows):
        row = tbl_meta.rows[idx]
        cell_lbl = row.cells[0]
        cell_val = row.cells[1]
        
        cell_lbl.width = Inches(2.0)
        cell_val.width = Inches(4.5)
        
        set_cell_background(cell_lbl, "F1F5F9")
        set_cell_margins(cell_lbl, 80, 80, 100, 100)
        set_cell_margins(cell_val, 80, 80, 100, 100)
        
        p0 = cell_lbl.paragraphs[0]
        p0.paragraph_format.space_before = Pt(0)
        p0.paragraph_format.space_after = Pt(0)
        r = p0.add_run(label)
        r.bold = True
        r.font.size = Pt(9)
        r.font.name = "Arial"
        
        p1 = cell_val.paragraphs[0]
        p1.paragraph_format.space_before = Pt(0)
        p1.paragraph_format.space_after = Pt(0)
        r = p1.add_run(val)
        r.font.size = Pt(9)
        r.font.name = "Arial"
        if idx == 5 and item.get("nilai") == "0.00":
            r.font.color.rgb = RGBColor(185, 28, 28)
            r.bold = True

    # Section 1: Checklist Bukti Dukung yang Wajib Diunggah
    h1 = doc.add_paragraph()
    h1.paragraph_format.space_before = Pt(14)
    h1.paragraph_format.space_after = Pt(6)
    r_h1 = h1.add_run("I. DAFTAR DOKUMEN BUKTI DUKUNG YANG WAJIB DISEDIAKAN")
    r_h1.bold = True
    r_h1.font.size = Pt(10)
    r_h1.font.name = "Arial"
    r_h1.font.color.rgb = RGBColor(30, 58, 138)

    bukti_list = item.get("bukti", [])
    if not bukti_list:
        p = doc.add_paragraph()
        r = p.add_run("• Tidak ada dokumen fisik khusus yang dipersyaratkan.")
        r.font.size = Pt(9)
        r.font.italic = True
    else:
        tbl_bukti = doc.add_table(rows=len(bukti_list) + 1, cols=4)
        tbl_bukti.alignment = WD_TABLE_ALIGNMENT.CENTER
        tbl_bukti.autofit = False
        
        # Headers
        headers = ["No", "Rincian Dokumen Bukti Dukung", "Format", "Status Checklist"]
        widths = [Inches(0.5), Inches(4.0), Inches(1.0), Inches(1.0)]
        
        for c_idx, h_text in enumerate(headers):
            c = tbl_bukti.rows[0].cells[c_idx]
            c.width = widths[c_idx]
            set_cell_background(c, "1E3A8A")
            set_cell_margins(c, 80, 80, 100, 100)
            p = c.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            r = p.add_run(h_text)
            r.bold = True
            r.font.size = Pt(8.5)
            r.font.color.rgb = RGBColor(255, 255, 255)
            r.font.name = "Arial"
            
        for b_idx, b_text in enumerate(bukti_list):
            row = tbl_bukti.rows[b_idx + 1]
            row.cells[0].width = widths[0]
            row.cells[1].width = widths[1]
            row.cells[2].width = widths[2]
            row.cells[3].width = widths[3]
            
            for cell in row.cells:
                set_cell_margins(cell, 80, 80, 100, 100)
                if (b_idx + 1) % 2 == 0:
                    set_cell_background(cell, "F8FAFC")
                    
            p0 = row.cells[0].paragraphs[0]
            p0.alignment = WD_ALIGN_PARAGRAPH.CENTER
            r0 = p0.add_run(str(b_idx + 1))
            r0.font.size = Pt(8.5)
            
            p1 = row.cells[1].paragraphs[0]
            r1 = p1.add_run(b_text)
            r1.font.size = Pt(8.5)
            r1.bold = True
            
            p2 = row.cells[2].paragraphs[0]
            p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
            fmt = "PDF / Word"
            if "SK" in b_text: fmt = "PDF (SK Resmi)"
            elif "Notula" in b_text or "Risalah" in b_text: fmt = "PDF/Word"
            elif "Foto" in b_text or "Banner" in b_text: fmt = "JPG/PDF"
            elif "Laporan" in b_text: fmt = "PDF (Laporan)"
            elif "Tautan" in b_text or "Screenshot" in b_text: fmt = "PDF (Tangkapan Layar)"
            r2 = p2.add_run(fmt)
            r2.font.size = Pt(8)
            
            p3 = row.cells[3].paragraphs[0]
            p3.alignment = WD_ALIGN_PARAGRAPH.CENTER
            r3 = p3.add_run("[   ] Siap")
            r3.font.size = Pt(8)

    # Note / Catatan Tindak Lanjut
    if item.get("catatan"):
        h_catatan = doc.add_paragraph()
        h_catatan.paragraph_format.space_before = Pt(12)
        h_catatan.paragraph_format.space_after = Pt(4)
        r_hc = h_catatan.add_run("⚠️ CATATAN & ATENSI KHUSUS TIM ZI BPS KEPULAUAN SELAYAR:")
        r_hc.bold = True
        r_hc.font.size = Pt(9.5)
        r_hc.font.name = "Arial"
        r_hc.font.color.rgb = RGBColor(185, 28, 28)
        
        p_c = doc.add_paragraph()
        p_c.paragraph_format.space_before = Pt(0)
        p_c.paragraph_format.space_after = Pt(8)
        r_c = p_c.add_run(item["catatan"])
        r_c.font.size = Pt(9)
        r_c.font.name = "Arial"
        r_c.font.color.rgb = RGBColor(153, 27, 27)

    # Section 2: Panduan & Batang Tubuh Dokumen
    h2 = doc.add_paragraph()
    h2.paragraph_format.space_before = Pt(14)
    h2.paragraph_format.space_after = Pt(6)
    r_h2 = h2.add_run("II. FORMAT DAN KERANGKA PENYUSUNAN DOKUMEN")
    r_h2.bold = True
    r_h2.font.size = Pt(10)
    r_h2.font.name = "Arial"
    r_h2.font.color.rgb = RGBColor(30, 58, 138)

    p_framework = doc.add_paragraph()
    p_framework.paragraph_format.space_before = Pt(0)
    p_framework.paragraph_format.space_after = Pt(8)
    r_f = p_framework.add_run(
        f"Dalam rangka memenuhi penilaian LKE indikator {item['code']}, dokumen bukti dukung yang disusun oleh tim kerja BPS Kabupaten Kepulauan Selayar wajib mencakup komponen substansi berikut:\n\n"
        f"1. Judul & Identitas Resmi:\n"
        f"   Dokumen diterbitkan dengan kop resmi Badan Pusat Statistik Kabupaten Kepulauan Selayar, bertanggal tahun berjalan (2025), serta memuat nomor naskah dinas resmi satker (Format nomor: .../73010/KU.000/.../2025).\n\n"
        f"2. Substansi Bukti Dukung:\n"
        f"   - Narasi latar belakang kebutuhan pelaksanaan kegiatan pembangunan ZI.\n"
        f"   - Matriks rencana aksi, target output, serta realisasi tindak lanjut yang dapat diverifikasi.\n"
        f"   - Lampiran data pendukung otentik (daftar hadir rapat, dokumentasi foto ber-geotagging/keterangan tanggal, tangkapan layar sistem internal BPS seperti SIMPEG, Sipandu, KipApp, atau BOS).\n\n"
        f"3. Validasi & Pengesahan:\n"
        f"   Dokumen ditandatangani oleh pejabat yang berwenang (Kepala BPS Kabupaten Kepulauan Selayar atau Ketua Tim Kerja ZI) dan dibubuhi cap dinas resmi."
    )
    r_f.font.size = Pt(9)
    r_f.font.name = "Arial"

    # Section 3: Tata Cara Penamaan & Penyimpanan di Google Drive
    h3 = doc.add_paragraph()
    h3.paragraph_format.space_before = Pt(12)
    h3.paragraph_format.space_after = Pt(6)
    r_h3 = h3.add_run("III. STANDAR PENAMAAN FILE & LOKASI GOOGLE DRIVE")
    r_h3.bold = True
    r_h3.font.size = Pt(10)
    r_h3.font.name = "Arial"
    r_h3.font.color.rgb = RGBColor(30, 58, 138)

    safe_q = "".join([c if c.isalnum() else "_" for c in item["question"][:30]])
    folder_name = item.get("drive_folder") or f"{item['code']}_{safe_q}"
    
    p_drive = doc.add_paragraph()
    p_drive.paragraph_format.space_before = Pt(0)
    p_drive.paragraph_format.space_after = Pt(16)
    r_d = p_drive.add_run(
        f"• Lokasi Folder Google Drive : /Bukti_Dukung_Selayar_2025/{area_num:02d}_{area_title.replace(' ', '_')}/{aspect}/{folder_name}/\n"
        f"• Standar Format Penamaan File : {item['code'].replace('.', '')}_[NamaDokumen]_BPS_Selayar_2025.pdf\n"
        f"• Hak Akses Folder Drive       : Wajib di-setting 'Anyone with the link can view' (Siapa saja yang memiliki tautan dapat melihat) agar dapat diakses oleh TPI dan TPN."
    )
    r_d.font.size = Pt(8.5)
    r_d.font.name = "Courier New"
    r_d.font.color.rgb = RGBColor(51, 65, 85)

    # Signature Block
    p_sign = doc.add_paragraph()
    p_sign.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p_sign.paragraph_format.space_before = Pt(12)
    p_sign.paragraph_format.space_after = Pt(2)
    r_s1 = p_sign.add_run("Benteng, Kepulauan Selayar, 2025\nMengetahui,\n")
    r_s1.font.size = Pt(9)
    r_s1.font.name = "Arial"
    
    r_s2 = p_sign.add_run("Kepala Badan Pusat Statistik\nKabupaten Kepulauan Selayar\n\n\n\n\n")
    r_s2.font.size = Pt(9)
    r_s2.bold = True
    r_s2.font.name = "Arial"
    
    r_s3 = p_sign.add_run("( .............................................................. )\n")
    r_s3.bold = True
    r_s3.font.size = Pt(9.5)
    r_s3.font.name = "Arial"
    
    r_s4 = p_sign.add_run("NIP. .................................................")
    r_s4.font.size = Pt(8.5)
    r_s4.font.name = "Arial"

    doc.save(filepath)

print("Generator function ready. Proceeding to build all folders and documents...")

# Mapping of specific pre-built templates to their exact indicator codes
TEMPLATE_MAP = {
    "1.a": "Format_01_SK_Tim_Pembangunan_ZI_2025_BPS_Selayar.docx",
    "1.b": "Format_02_Notula_Rapat_Pembentukan_Tim_ZI_BPS_Selayar.docx",
    "3.b": "Format_04_Laporan_Monev_ZI_Triwulan_BPS_Selayar.docx",
    "4.a": "Format_07_Laporan_Penerapan_Role_Model_Pimpinan_BPS_Selayar.docx",
    "4.b": "Format_08_Laporan_Agen_Perubahan_BPS_Selayar.docx",
    "1.c": "Format_09_Kertas_Kerja_Evaluasi_SOP_BPS_Selayar.docx", # In Area 2
    "5.c": "Format_03_Surat_Pernyataan_Bebas_Benturan_Kepentingan_BPS_Selayar.docx", # In Area 5
    "3.a": "Format_10_Laporan_Pengelolaan_Pengaduan_dan_WBS_BPS_Selayar.docx", # In Area 5
    "2.c": "Format_06_Pedoman_Reward_Punishment_Petugas_Layanan_BPS_Selayar.docx", # In Area 6
}

created_folders = 0
created_docs = 0

for target_root in targets:
    print(f"\nProcessing target root: {target_root}")
    os.makedirs(target_root, exist_ok=True)
    
    # Copy Master Checklist Excel to root
    excel_src = "templates/Checklist_Bukti_Dukung_ZI_BPS_Kepulauan_Selayar_2025.xlsx"
    if os.path.exists(excel_src):
        shutil.copy2(excel_src, os.path.join(target_root, "00_MASTER_CHECKLIST_LKE_ZI_2025_BPS_SELAYAR.xlsx"))
        # Also copy the PDF guide
        if os.path.exists("04. Pemetaan bukti dukung LKE Pembangunan ZI 2025.pdf"):
            shutil.copy2("04. Pemetaan bukti dukung LKE Pembangunan ZI 2025.pdf", os.path.join(target_root, "00_PANDUAN_RESMI_PEMETAAN_BUKTI_DUKUNG_BPS_RI.pdf"))

    for area in app_data["areas"]:
        area_num = area["number"]
        area_title = area["title"]
        area_dir_name = f"{area_num:02d}_{sanitize_folder_name(area_title)}"
        
        # Process Pemenuhan & Reform
        for aspect_key, aspect_label in [("pemenuhan", "Pemenuhan"), ("reform", "Reform")]:
            aspect_data = area[aspect_key]
            for grp in aspect_data["groups"]:
                grp_name = grp["name"]
                for itm in grp["items"]:
                    code = itm["code"]
                    clean_code = code.replace(".", "_")
                    drive_fld = itm.get("drive_folder")
                    if not drive_fld:
                        clean_q = sanitize_folder_name(itm["question"][:40])
                        drive_fld = f"{clean_code}_{clean_q}"
                    else:
                        drive_fld = sanitize_folder_name(drive_fld)
                        
                    folder_path = os.path.join(target_root, area_dir_name, aspect_label, drive_fld)
                    os.makedirs(folder_path, exist_ok=True)
                    created_folders += 1
                    
                    # Generate official docs version for this indicator
                    doc_filename = f"Panduan_Format_{clean_code}_{drive_fld[:35]}.docx"
                    doc_path = os.path.join(folder_path, doc_filename)
                    create_indicator_docx(doc_path, itm, area_num, area_title, aspect_label, grp_name)
                    created_docs += 1
                    
                    # If specific template matches this indicator, copy it into this folder too
                    # Check Area specifics
                    if area_num == 1 and code in ["1.a", "1.b", "3.b", "4.a", "4.b"]:
                        tpl_name = TEMPLATE_MAP.get(code)
                        if tpl_name and os.path.exists(f"templates/{tpl_name}"):
                            shutil.copy2(f"templates/{tpl_name}", os.path.join(folder_path, tpl_name))
                    elif area_num == 2 and code == "1.c":
                        tpl_name = TEMPLATE_MAP.get(code)
                        if tpl_name and os.path.exists(f"templates/{tpl_name}"):
                            shutil.copy2(f"templates/{tpl_name}", os.path.join(folder_path, tpl_name))
                    elif area_num == 5 and code in ["5.c", "3.a"]:
                        tpl_name = TEMPLATE_MAP.get(code)
                        if tpl_name and os.path.exists(f"templates/{tpl_name}"):
                            shutil.copy2(f"templates/{tpl_name}", os.path.join(folder_path, tpl_name))
                    elif area_num == 6 and code == "2.c":
                        tpl_name = TEMPLATE_MAP.get(code)
                        if tpl_name and os.path.exists(f"templates/{tpl_name}"):
                            shutil.copy2(f"templates/{tpl_name}", os.path.join(folder_path, tpl_name))

print(f"\nSUCCESS! Generated {created_folders} indicator folders and {created_docs} formatted Word documents (.docx) across all target locations!")

