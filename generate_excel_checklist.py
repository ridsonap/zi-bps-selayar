import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
import json
import re

wb = openpyxl.Workbook()
wb.remove(wb.active)

fill_header = PatternFill(start_color="1A365D", end_color="1A365D", fill_type="solid")
font_header = Font(name="Calibri", size=11, bold=True, color="FFFFFF")

fill_success = PatternFill(start_color="DCFCE7", end_color="DCFCE7", fill_type="solid")
font_success = Font(name="Calibri", size=10, bold=True, color="166534")

fill_warning = PatternFill(start_color="FEF3C7", end_color="FEF3C7", fill_type="solid")
font_warning = Font(name="Calibri", size=10, bold=True, color="92400E")

fill_danger = PatternFill(start_color="FEE2E2", end_color="FEE2E2", fill_type="solid")
font_danger = Font(name="Calibri", size=10, bold=True, color="991B1B")

thin_border = Border(
    left=Side(style='thin', color='CBD5E1'),
    right=Side(style='thin', color='CBD5E1'),
    top=Side(style='thin', color='CBD5E1'),
    bottom=Side(style='thin', color='CBD5E1')
)

with open("index.html", "r", encoding="utf-8") as f:
    html_text = f.read()

# Extract APP_DATA
match = re.search(r'const APP_DATA\s*=\s*(\{.*?\});\s*\n', html_text)
if not match:
    # try broader search
    idx_start = html_text.find('const APP_DATA = {')
    idx_end = html_text.find('};\n', idx_start)
    app_data = json.loads(html_text[idx_start + len('const APP_DATA = '):idx_end + 1])
else:
    app_data = json.loads(match.group(1))

# Extract items
items = []
for area in app_data.get("areas", []):
    for grp in area.get("pemenuhan", {}).get("groups", []):
        for itm in grp.get("items", []):
            items.append({
                "areaNum": area.get("number"),
                "areaTitle": area.get("title"),
                "aspect": "Pemenuhan",
                "code": itm.get("code"),
                "question": itm.get("question"),
                "bukti": itm.get("bukti", []),
                "nilai": itm.get("nilai"),
                "status": itm.get("badge", itm.get("status")),
                "jawaban": itm.get("jawaban"),
                "narasi": itm.get("narasi", "")
            })
    for grp in area.get("reform", {}).get("groups", []):
        for itm in grp.get("items", []):
            items.append({
                "areaNum": area.get("number"),
                "areaTitle": area.get("title"),
                "aspect": "Reform",
                "code": itm.get("code"),
                "question": itm.get("question"),
                "bukti": itm.get("bukti", []),
                "nilai": itm.get("nilai"),
                "status": itm.get("badge", itm.get("status")),
                "jawaban": itm.get("jawaban"),
                "narasi": itm.get("narasi", "")
            })

# 1. Summary Sheet
ws_sum = wb.create_sheet(title="Ringkasan Nilai LKE")
ws_sum.views.sheetView[0].showGridLines = True

ws_sum.merge_cells("A1:G1")
ws_sum["A1"] = "LEMBAR KERJA EVALUASI (LKE) PEMBANGUNAN ZONA INTEGRITAS MENUJU WBK/WBBM"
ws_sum["A1"].font = Font(name="Calibri", size=14, bold=True, color="1A365D")
ws_sum["A1"].alignment = Alignment(horizontal="center", vertical="center")

ws_sum.merge_cells("A2:G2")
ws_sum["A2"] = "BADAN PUSAT STATISTIK KABUPATEN KEPULAUAN SELAYAR — TAHUN 2025"
ws_sum["A2"].font = Font(name="Calibri", size=11, bold=True, color="C6953B")
ws_sum["A2"].alignment = Alignment(horizontal="center", vertical="center")

headers_sum = ["No", "Komponen / Area Perubahan", "Bobot", "Capaian Pemenuhan", "Capaian Reform", "Total Nilai", "% Capaian"]
for col_idx, h in enumerate(headers_sum, 1):
    cell = ws_sum.cell(row=4, column=col_idx, value=h)
    cell.fill = fill_header
    cell.font = font_header
    cell.alignment = Alignment(horizontal="center", vertical="center")

summary_data = [
    ("1", "Manajemen Perubahan", 8.00, 2.54, 4.00, 6.54, "81.75%"),
    ("2", "Penataan Tatalaksana", 7.00, 3.38, 3.50, 6.88, "98.29%"),
    ("3", "Penataan Sistem Manajemen SDM Aparatur", 10.00, 3.96, 3.50, 7.46, "74.60%"),
    ("4", "Penguatan Akuntabilitas Kinerja", 10.00, 5.00, 3.50, 8.50, "85.00%"),
    ("5", "Penguatan Pengawasan", 15.00, 3.38, 7.50, 10.88, "72.53%"),
    ("6", "Peningkatan Kualitas Pelayanan Publik", 10.00, 3.95, 2.50, 6.45, "64.50%"),
    ("A", "TOTAL PENGUNGKIT", 60.00, 22.20, 24.50, 46.70, "77.83%"),
    ("B", "KOMPONEN HASIL (Menunggu Survei Eksternal)", 40.00, 0.00, 0.00, 0.00, "0.00%"),
    ("", "TOTAL NILAI ZONA INTEGRITAS", 100.00, 22.20, 24.50, 46.70, "46.70%")
]

for row_idx, r_data in enumerate(summary_data, 5):
    for col_idx, val in enumerate(r_data, 1):
        cell = ws_sum.cell(row=row_idx, column=col_idx, value=val)
        cell.alignment = Alignment(horizontal="left" if col_idx == 2 else "center", vertical="center")
        cell.border = thin_border
        if row_idx in [11, 12, 13]:
            cell.font = Font(name="Calibri", size=10, bold=True)
            cell.fill = PatternFill(start_color="F1F5F9", end_color="F1F5F9", fill_type="solid")

ws_sum.column_dimensions["A"].width = 6
ws_sum.column_dimensions["B"].width = 45
ws_sum.column_dimensions["C"].width = 12
ws_sum.column_dimensions["D"].width = 18
ws_sum.column_dimensions["E"].width = 16
ws_sum.column_dimensions["F"].width = 14
ws_sum.column_dimensions["G"].width = 14

# 2. Checklist Sheet
ws_chk = wb.create_sheet(title="Checklist Bukti Dukung")
ws_chk.views.sheetView[0].showGridLines = True

ws_chk.merge_cells("A1:H1")
ws_chk["A1"] = "CHECKLIST PEMENUHAN BUKTI DUKUNG LKE ZONA INTEGRITAS"
ws_chk["A1"].font = Font(name="Calibri", size=13, bold=True, color="1A365D")
ws_chk["A1"].alignment = Alignment(horizontal="center", vertical="center")

ws_chk.merge_cells("A2:H2")
ws_chk["A2"] = "BPS KABUPATEN KEPULAUAN SELAYAR — TAHUN 2025"
ws_chk["A2"].font = Font(name="Calibri", size=10, bold=True, color="C6953B")
ws_chk["A2"].alignment = Alignment(horizontal="center", vertical="center")

headers_chk = ["No", "Area", "Aspek", "Kode", "Pertanyaan Indikator", "Daftar Bukti Dukung yang Diperlukan", "Nilai", "Status LKE"]
for col_idx, h in enumerate(headers_chk, 1):
    cell = ws_chk.cell(row=4, column=col_idx, value=h)
    cell.fill = fill_header
    cell.font = font_header
    cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

for idx, itm in enumerate(items, 1):
    row_num = idx + 4
    ws_chk.cell(row=row_num, column=1, value=idx).alignment = Alignment(horizontal="center", vertical="top")
    ws_chk.cell(row=row_num, column=2, value=f"Area {itm.get('areaNum')}: {itm.get('areaTitle')}").alignment = Alignment(horizontal="left", vertical="top")
    ws_chk.cell(row=row_num, column=3, value=itm.get('aspect')).alignment = Alignment(horizontal="center", vertical="top")
    ws_chk.cell(row=row_num, column=4, value=itm.get('code')).alignment = Alignment(horizontal="center", vertical="top")
    
    q_cell = ws_chk.cell(row=row_num, column=5, value=itm.get('question'))
    q_cell.alignment = Alignment(horizontal="left", vertical="top", wrap_text=True)
    
    bukti_list = itm.get('bukti', [])
    if isinstance(bukti_list, list):
        bukti_text = "\n• ".join([""] + bukti_list).strip()
    else:
        bukti_text = str(bukti_list)
    b_cell = ws_chk.cell(row=row_num, column=6, value=bukti_text)
    b_cell.alignment = Alignment(horizontal="left", vertical="top", wrap_text=True)
    
    n_cell = ws_chk.cell(row=row_num, column=7, value=itm.get('nilai'))
    n_cell.alignment = Alignment(horizontal="center", vertical="top")
    
    status_label = itm.get('status')
    st_cell = ws_chk.cell(row=row_num, column=8, value=status_label)
    st_cell.alignment = Alignment(horizontal="center", vertical="top")
    
    if status_label in ['Lengkap', 'complete']:
        st_cell.fill = fill_success
        st_cell.font = font_success
        st_cell.value = "Lengkap"
    elif status_label in ['Perlu Dilengkapi', 'missing']:
        st_cell.fill = fill_danger
        st_cell.font = font_danger
        st_cell.value = "Perlu Dilengkapi"
    else:
        st_cell.fill = fill_warning
        st_cell.font = font_warning
        st_cell.value = "Sebagian"
        
    for c in range(1, 9):
        ws_chk.cell(row=row_num, column=c).border = thin_border

ws_chk.column_dimensions["A"].width = 5
ws_chk.column_dimensions["B"].width = 24
ws_chk.column_dimensions["C"].width = 12
ws_chk.column_dimensions["D"].width = 10
ws_chk.column_dimensions["E"].width = 40
ws_chk.column_dimensions["F"].width = 48
ws_chk.column_dimensions["G"].width = 8
ws_chk.column_dimensions["H"].width = 16

wb.save("templates/Checklist_Bukti_Dukung_ZI_BPS_Kepulauan_Selayar_2025.xlsx")
print(f"Checklist excel successfully created with {len(items)} items.")
